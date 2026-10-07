# Gera planilha/modelo-diagnostico-combo-vitalicio.xlsx a partir dos cabeçalhos do Code.gs e do catálogo real (src/catalogo.js).
import json, re, subprocess, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
root = pathlib.Path(__file__).resolve().parent.parent
gs = (root / 'apps-script' / 'Code.gs').read_text()
headers = json.loads(re.search(r'const HEADERS = (\{.*?\});\n// HEADERS:END', gs, re.S).group(1))
demo = json.loads(subprocess.check_output(['node', str(root / 'src' / 'catalogo.js')]))
version = subprocess.check_output(['node', '-p', "require('./src/catalogo').VERSION"], cwd=root).decode().strip()
wb = Workbook(); wb.remove(wb.active)
bold = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='23655C')
readme = wb.create_sheet('Leia-me')
for line in [
    ['Modelo da planilha do Diagnóstico do Combo Vitalício'],
    [''],
    ['Abas preenchidas automaticamente pelo diagnóstico (não editar à mão): Diagnósticos, Planos de estudo, Cronograma, Sugestões de atualização.'],
    ['Abas mantidas pela equipe: Catálogo (uma linha por aula) e Configuração (versão do catálogo).'],
    ['ATENÇÃO: módulos e aulas vêm do documento da equipe (07/10/2026), mas SEM duração e link. Temas, nível e pré-requisitos são uma PROPOSTA a validar. Preencha duracao_min e link_aula.'],
    ['Enquanto um curso não tiver aulas, o plano usa blocos de estudo sobre carga_estimada_h. Ao cadastrar as aulas: uma linha por aula, repetindo os dados do curso.'],
    ['tipo: "curso" entra no plano; "recurso" (ebook, aulas ao vivo, bônus, suporte) não entra no cronograma.'],
    ['Catálogo: separe vários temas ou pré-requisitos com ponto e vírgula. Nível: 1 Iniciante, 2 Básico, 3 Intermediário, 4 Avançado.'],
    ['Catálogo: duração em minutos (número). Deixe vazio se não souber: o plano usa uma estimativa sinalizada ao aluno.'],
    ['Catálogo: links devem começar com https://. Para tirar uma aula do ar sem apagar, coloque "não" em ativo.'],
    ['Ao alterar o catálogo, atualize versao_catalogo na aba Configuração.'],
    ['Compartilhamento: restrinja a planilha à equipe autorizada. Ela contém nome e e-mail de alunos.'],
]:
    readme.append(line)
readme['A1'].font = Font(bold=True, size=14); readme.column_dimensions['A'].width = 140
for name, cols in headers.items():
    ws = wb.create_sheet(name)
    ws.append(cols)
    for c in ws[1]: c.font = bold; c.fill = fill; c.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'A2'
    for i, col in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(40, len(col) + 6))
    if name == 'Catálogo':
        for r in demo: ws.append([r.get(c, '') for c in cols])
        dv = DataValidation(type='list', formula1='"sim,não"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'R2:R1000')
        dv2 = DataValidation(type='whole', operator='between', formula1='1', formula2='4'); ws.add_data_validation(dv2); dv2.add('E2:E1000')
    if name == 'Configuração':
        ws.append(['versao_catalogo', version, 'Atualize sempre que o catálogo mudar. Fica registrada em cada plano.'])
out = root / 'planilha' / 'modelo-diagnostico-combo-vitalicio.xlsx'
wb.save(out); print('ok', out)
